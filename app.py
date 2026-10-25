import os
import re
from collections import Counter
from flask import Flask, render_template, request
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024  # 5 MB upload limit
ALLOWED_EXTENSIONS = {"pdf", "docx"}

ROLE_KEYWORDS = {
    "Data Analyst": ["python", "sql", "excel", "power bi", "tableau", "statistics", "data analysis", "pandas", "visualization"],
    "Web Developer": ["html", "css", "javascript", "react", "flask", "api", "git", "responsive design", "database"],
    "AI Engineer": ["python", "machine learning", "deep learning", "nlp", "pandas", "numpy", "scikit-learn", "tensorflow", "data preprocessing"],
    "Cloud Engineer": ["cloud", "aws", "azure", "gcp", "linux", "docker", "kubernetes", "networking", "ci/cd"]
}

def extract_text(file_storage):
    """Extract text from an uploaded PDF or DOCX file."""
    filename = secure_filename(file_storage.filename or "")
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Please upload a PDF or DOCX file.")
    if extension == "pdf":
        try:
            import fitz
        except ImportError:
            raise ValueError("PDF support is missing. Run: pip install PyMuPDF")
        document = fitz.open(stream=file_storage.read(), filetype="pdf")
        text = "\n".join(page.get_text() for page in document)
        document.close()
    else:
        try:
            from docx import Document
        except ImportError:
            raise ValueError("DOCX support is missing. Run: pip install python-docx")
        document = Document(file_storage)
        text = "\n".join(p.text for p in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                text += "\n" + " ".join(cell.text for cell in row.cells)
    if not text.strip():
        raise ValueError("No readable text was found. Try a text-based PDF or DOCX file.")
    return text

def analyze_resume(text, role):
    """A transparent, rule-based prototype; scores are guidance, not an official ATS result."""
    lower = text.lower()
    sections = {
        "Contact information": any(x in lower for x in ["@", "phone", "mobile", "+91"]),
        "Education": any(x in lower for x in ["education", "b.tech", "bachelor", "degree", "university", "college"]),
        "Skills": any(x in lower for x in ["skills", "technical skills", "technologies"]),
        "Projects": "project" in lower,
        "Experience or internship": any(x in lower for x in ["experience", "internship", "intern"]),
        "Certifications": any(x in lower for x in ["certification", "certifications", "certificate"])
    }
    section_score = round(sum(sections.values()) / len(sections) * 40)
    words = re.findall(r"\b[\w+#.]+\b", lower)
    length_score = 20 if 150 <= len(words) <= 900 else (12 if 80 <= len(words) < 150 or 900 < len(words) <= 1300 else 5)
    action_words = ["built", "developed", "created", "designed", "implemented", "analyzed", "improved", "managed", "led"]
    action_score = min(15, sum(lower.count(word) for word in action_words) * 3)
    contact_score = 10 if "@" in lower else 0
    keyword_list = ROLE_KEYWORDS[role]
    matched = [word for word in keyword_list if word in lower]
    missing = [word for word in keyword_list if word not in lower]
    keyword_score = round(len(matched) / len(keyword_list) * 15)
    total = min(100, section_score + length_score + action_score + contact_score + keyword_score)

    suggestions = []
    if not sections["Contact information"]:
        suggestions.append("Add clear contact information, including a professional email address.")
    if not sections["Education"]:
        suggestions.append("Add an Education section with your degree, institution, and expected graduation date.")
    if not sections["Skills"]:
        suggestions.append("Add a clearly labelled Skills or Technical Skills section.")
    if not sections["Projects"]:
        suggestions.append("Include 2–3 relevant projects and briefly describe your contribution.")
    if not sections["Experience or internship"]:
        suggestions.append("If applicable, add internships, volunteering, or practical experience. Do not invent experience.")
    if not sections["Certifications"]:
        suggestions.append("Add relevant certifications if you have them; omit this section if you do not.")
    if missing:
        suggestions.append("If you genuinely have these skills, consider adding relevant evidence for: " + ", ".join(missing[:6]) + ".")
    if not suggestions:
        suggestions.append("Your resume contains the main sections checked by this prototype. Tailor project descriptions to the selected role.")
    if not any(w in lower for w in ["improved", "reduced", "increased", "achieved", "%"]):
        suggestions.append("Where accurate, describe outcomes with measurable results (for example, time saved or accuracy improved).")
    return {
        "score": total, "sections": sections, "matched": matched, "missing": missing,
        "keyword_score": keyword_score, "word_count": len(words), "suggestions": suggestions,
        "role": role, "section_score": section_score
    }

@app.route("/", methods=["GET", "POST"])
def index():
    result = None
    error = None
    role = request.form.get("role", "AI Engineer")
    if request.method == "POST":
        uploaded = request.files.get("resume")
        if not uploaded or not uploaded.filename:
            error = "Choose a resume file first."
        else:
            try:
                text = extract_text(uploaded)
                result = analyze_resume(text, role if role in ROLE_KEYWORDS else "AI Engineer")
            except Exception as exc:
                error = str(exc)
    return render_template("index.html", result=result, error=error, roles=list(ROLE_KEYWORDS.keys()), selected_role=role)

@app.errorhandler(413)
def too_large(_error):
    return render_template("index.html", result=None, error="The file is too large. Please upload a file smaller than 5 MB.", roles=list(ROLE_KEYWORDS.keys()), selected_role="AI Engineer"), 413

if __name__ == "__main__":
    app.run(debug=True)
